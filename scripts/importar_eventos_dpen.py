"""Gera a migration com os eventos anteriores e as vendas por dia da planilha DPEN 25-26 (Heitor, 08/10/2026).

Uso: python3 -I scripts/importar_eventos_dpen.py <planilha.xlsx> > supabase/migrations/0037_eventos_anteriores_dpen.sql
- Cada linha de "Resumo Eventos" vira um evento finalizado (nome = apelido; nome do DPEN guardado em dpen_nome).
- Dias de funcionamento vêm de "Eventos por Dia(DB)"; abertura do 1º dia e fechamento do último vêm do resumo.
- Vendas por dia e produto vêm de "Produtos por Dia(DB)" (as duas contas somadas).
- Produto vendido que ainda não tem ficha vira um produto sem ficha (versão 0), para entrar nas comparações.
- O cardápio de cada evento é o que vendeu, com o preço médio praticado.
"""
import collections, datetime, sys
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
num = lambda v: v if isinstance(v, (int, float)) and not isinstance(v, bool) else None
q = lambda s: 'null' if s is None else "'" + str(s).replace("'", "''") + "'"
n = lambda v: 'null' if v is None else repr(round(float(v), 2))
MESES = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro', 'outubro', 'novembro', 'dezembro']
PEQUENAS = {'de', 'da', 'do', 'das', 'dos', 'na', 'no', 'com', 'e', 'ao', 'a', 'o', 'em'}


def bonito(s):
    """'FOCA NO PARMA' -> 'Foca no Parma'; tamanhos 'G'/'P' ficam maiúsculos."""
    p = s.strip().lower().split()
    return ' '.join(w.upper() if w in ('g', 'p') else w if (i and w in PEQUENAS) else w[:1].upper() + w[1:] for i, w in enumerate(p))


obs = lambda e: 'Importado da planilha DPEN 25-26 em 08/10/2026. Faturamento no resumo: R$ ' + f"{e['total']:,.0f}".replace(',', '.') + '.'
valido = lambda v: None if v in (None, '', 'Inválido') else v

# Eventos
resumo = [r for r in list(wb['Resumo Eventos'].iter_rows(values_only=True))[3:] if r[0] and isinstance(r[6], datetime.datetime)]
apelidos = collections.Counter(r[1] for r in resumo)
eventos = []
for r in resumo:
    nome = r[1].strip()
    if apelidos[r[1]] > 1:
        nome += f' ({MESES[r[6].month - 1]})'
    eventos.append({'dpen': r[0].strip(), 'nome': nome, 'cidade': valido(r[2]), 'local': valido(r[3]), 'gastronomia': valido(r[4]),
                    'barracas': num(r[5]), 'inicio': r[6], 'fim': r[7] if isinstance(r[7], datetime.datetime) else None,
                    'organizador': 'Cola em Sampa' if 'COLA EM SAMPA' in r[0].upper() else None, 'total': num(r[11])})
por_nome = {e['dpen']: e for e in eventos}

dias = collections.defaultdict(set)
for r in list(wb['Eventos por Dia(DB)'].iter_rows(values_only=True))[3:]:
    if r[2] and isinstance(r[4], datetime.datetime):
        assert r[2].strip() in por_nome, r[2]
        dias[r[2].strip()].add(r[4].date())

# Vendas: (evento, dia, produto) -> [quantidade, total]
vendas = collections.defaultdict(lambda: [0.0, 0.0])
ultimo = {}
for r in list(wb['Produtos por Dia(DB)'].iter_rows(values_only=True))[3:]:
    if not r[2] or not r[5] or not isinstance(r[3], datetime.datetime):
        continue
    ev = r[2].strip()
    assert ev in por_nome, ev
    p = r[5].strip().upper()
    v = vendas[(ev, r[3].date(), p)]
    v[0] += num(r[7]) or 0
    v[1] += num(r[8]) or 0
    if num(r[6]) is not None and (p not in ultimo or r[3] > ultimo[p][0]):
        ultimo[p] = (r[3], r[6])
    dias[ev].add(r[3].date())

linha_status = {}
for r in list(wb['Resumo Produtos'].iter_rows(values_only=True))[1:]:
    if r[1]:
        linha_status[r[1].strip().upper()] = (r[3], r[2])

out = ['-- Eventos anteriores (jan/2025 a ago/2026) e vendas por dia, da planilha DPEN 25-26 enviada pelo Heitor em 08/10.',
       '-- Gerado por scripts/importar_eventos_dpen.py. Só insere: nada do que já existe muda.', '']

# Produtos vendidos que ainda não têm ficha
prods = sorted({k[2] for k in vendas})
out.append('insert into receitas (nome, tipo, linha, origem, unidade, preco_venda, ativo, versao_atual)')
out.append('select x.nome, \'produto\', x.linha, \'propria\', \'un\', x.preco, x.ativo, 0 from (values')
linhas = []
for p in prods:
    linha, status = linha_status.get(p, (None, 'ATIVO'))
    linhas.append(f'  ({q(bonito(p))}, {q(bonito(linha) if linha and linha != "GERAL" else None)}, {n(ultimo.get(p, (None, None))[1])}, {"false" if status == "INATIVO" else "true"})')
out.append(',\n'.join(linhas))
out.append(') as x (nome, linha, preco, ativo)')
out.append("where not exists (select 1 from receitas r where lower(r.nome) = lower(x.nome) and r.tipo = 'produto');")
out.append('')

out.append('insert into eventos (nome, status, tipo, organizador, cidade, local, gastronomia, barracas, dpen_nome, observacao) values')
out.append(',\n'.join(
    f"  ({q(e['nome'])}, 'finalizado', 'Festival gastronômico', {q(e['organizador'])}, {q(e['cidade'])}, {q(e['local'])}, {q(e['gastronomia'])}, "
    f"{n(e['barracas'])}, {q(e['dpen'])}, {q(obs(e))})"
    for e in eventos) + ';')
out.append('')

out.append('insert into evento_dias (evento_id, data, abre, fecha)')
out.append('select e.id, x.d::date, x.a::time, x.f::time from (values')
ds = []
for e in eventos:
    ld = sorted(dias[e['dpen']])
    for i, d in enumerate(ld):
        abre = e['inicio'].strftime('%H:%M') if i == 0 else None
        fecha = e['fim'].strftime('%H:%M') if (i == len(ld) - 1 and e['fim'] and e['fim'].strftime('%H:%M') != '00:00') else None
        ds.append(f"  ({q(e['dpen'])}, '{d}', {q(abre)}, {q(fecha)})")
out.append(',\n'.join(ds))
out.append(') as x (n, d, a, f) join eventos e on e.dpen_nome = x.n;')
out.append('')

out.append("insert into evento_vendas (evento_id, data, receita_id, produto, quantidade, total, origem)")
out.append("select e.id, x.d::date, r.id, r.nome, x.qt, x.tt, 'dpen' from (values")
out.append(',\n'.join(f"  ({q(ev)}, '{d}', {q(bonito(p))}, {n(v[0])}, {n(v[1])})" for (ev, d, p), v in sorted(vendas.items())))
out.append(") as x (n, d, p, qt, tt) join eventos e on e.dpen_nome = x.n join receitas r on lower(r.nome) = lower(x.p) and r.tipo = 'produto';")
out.append('')

# Cardápio: o que vendeu em cada evento, com o preço médio praticado, do mais vendido para o menos vendido.
card = collections.defaultdict(lambda: [0.0, 0.0])
for (ev, d, p), v in vendas.items():
    card[(ev, p)][0] += v[0]
    card[(ev, p)][1] += v[1]
cs = []
for ev in por_nome:
    itens = sorted(((p, v) for (e, p), v in card.items() if e == ev), key=lambda t: -t[1][0])
    for o, (p, v) in enumerate(itens, 1):
        cs.append(f"  ({q(ev)}, {q(bonito(p))}, {n(v[1] / v[0]) if v[0] else 'null'}, {o})")
out.append('insert into evento_produtos (evento_id, receita_id, preco, ordem)')
out.append('select e.id, r.id, x.pr, x.o from (values')
out.append(',\n'.join(cs))
out.append(") as x (n, p, pr, o) join eventos e on e.dpen_nome = x.n join receitas r on lower(r.nome) = lower(x.p) and r.tipo = 'produto';")

print('\n'.join(out))
print(f'-- {len(eventos)} eventos, {len(ds)} dias, {len(vendas)} vendas (dia × produto), {len(cs)} itens de cardápio, {len(prods)} produtos vendidos.', file=sys.stderr)
