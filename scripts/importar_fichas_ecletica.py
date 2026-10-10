"""Gera a migration 0070 com as fichas técnicas da Eclética (Heitor, 10/10) e o relatório de diferenças.

Uso: python3 -I scripts/importar_fichas_ecletica.py <fts.xlsx> <"Fichas Técnicas.xlsx"> <produtos.xlsx> <relatorio.md> \
       > supabase/migrations/0070_dados_fichas.sql

- fts.xlsx (Relatório de fórmulas da Eclética) é a base: cada fórmula vira uma ficha com rendimento 1 (1 kg, 1 L ou
  1 unidade do item). Código de material = pré-preparo (ligado ao item de estoque); código de produto = produto da loja.
  Item que tem fórmula própria entra como pré-preparo (sub-ficha), o resto como insumo pelo código.
- Pré-preparo que já existe nos eventos com o mesmo nome é a mesma ficha: ganha a versão da Eclética.
- "Fichas Técnicas.xlsx" (as fichas impressas de hoje) só dá o texto: modo de preparo, armazenamento, validade,
  observações, responsável e a porção (potinhos, seringas). As quantidades de lá vão para o relatório, para conferir.
"""
import re, sys, unicodedata
import openpyxl

sys.path.insert(0, __file__.rsplit('/', 1)[0])
from importar_cadastros import nome_bom, simples, sql, valores  # noqa: E402

LABELS = {'cód preparo', 'insumos', '#', 'rendimento', 'qtd em potinhos', 'modo de preparo', 'armazenamento', 'validade',
          'observações', 'obs', 'ajuste', 'ajuste 1', 'atualizada', 'porçoes', 'seringas', 'responsável', 'item'}
PORCAO = {'qtd em potinhos': 'potinhos', 'porçoes': 'porções', 'seringas': 'seringas'}
# Nome do pré-preparo nos eventos -> código da Eclética, quando o nome não é igual.
ALIAS = {'focaccia em forma': 100478, 'base amenteigada de cheesecake': 100497, 'mix de cogumelos': 100543,
         'calda de frutas vermelhas caseira': 100495, 'massa de cheesecake': 100498}
# Ficha impressa -> código, quando o título não bate com o nome da Eclética (ou o código impresso está errado).
IMPRESSA = {'molho de cheddar vigor': 100232, 'bacon triturado': None, 'pulled pork 75g': 100249, 'brownie 90g': 100385,
            'brigadeiro de pistache': None, 'disco de mussarela triturado': None, 'mix de cogumelos trufado': 100543,
            'focaccia forma': 100478, 'potinho de maionese verde': 100208, 'cookie recheado 120g': None}
FORA = re.compile(r'substituir|^$', re.I)
INATIVA = re.compile(r'n[ãa]o lan[çc]ar|futuro|pr[óo]ximo milkshake', re.I)


def low(v):
    return v.strip().lower() if isinstance(v, str) else None


def ler_fts(path):
    fichas, cur = [], None
    for r in openpyxl.load_workbook(path, read_only=True, data_only=True).active.iter_rows(values_only=True):
        if isinstance(r[0], str) and r[1] == 'Código':
            cod, nome = r[0].split(' - ', 1)
            cur = {'cod': int(cod), 'nome': ' '.join(nome.split()), 'itens': []}
            fichas.append(cur)
        elif cur and isinstance(r[0], int) and isinstance(r[2], (int, float)) and r[2] > 0:
            perda = r[6] if isinstance(r[6], (int, float)) else 0
            cur['itens'].append({'cod': r[0], 'nome': r[1], 'qtd': float(r[2]), 'unid': r[3], 'aprov': round(1 - perda / 100, 4),
                                 'delivery': r[7] == '*'})
    return [f for f in fichas if not FORA.search(f['nome'])]


def ler_impressas(path):
    out = []
    wb = openpyxl.load_workbook(path, data_only=True)
    for s in wb.sheetnames[1:]:
        rows = [list(r) for r in wb[s].iter_rows(values_only=True)]
        cur, modo = None, None
        for i, r in enumerate(rows):
            a = r[0] if r else None
            la = low(a)
            prox = [low(x[0]) for x in rows[i + 1:i + 3] if x and x[0] not in (None, '')]
            if isinstance(a, str) and la not in LABELS and prox and prox[0] in ('cód preparo', 'insumos'):
                cur = {'aba': s, 'titulo': ' '.join(a.split()), 'codigo': None, 'resp': None, 'itens': [], 'rend': None,
                       'porcoes': None, 'porcao_nome': None, 'modo': [], 'armaz': None, 'validade': None, 'obs': []}
                out.append(cur)
                modo = None
                continue
            if not cur:
                continue
            vals = [v for v in r[1:] if v not in (None, '')]
            vals = [v for v in vals if low(v) not in LABELS]
            if la == 'cód preparo':
                cur['codigo'] = r[1] if isinstance(r[1], int) else None
                for k, v in enumerate(r):
                    if low(v) == 'responsável':
                        resto = [x for x in r[k + 1:] if isinstance(x, str) and x.strip() and low(x) not in LABELS]
                        if resto:
                            cur['resp'] = resto[0].strip().title()
                modo = None
            elif la == 'rendimento':
                cur['rend'] = r[2] if len(r) > 2 else None
                modo = None
            elif la in PORCAO:
                cur['porcoes'], cur['porcao_nome'], modo = (r[2] if len(r) > 2 else None), PORCAO[la], None
            elif la == 'modo de preparo':
                modo = 'modo'
            elif la == 'armazenamento':
                if vals and not cur['armaz']:
                    cur['armaz'] = str(vals[0]).strip()
                modo = 'armaz'
            elif la == 'validade':
                if vals:
                    cur['validade'] = str(vals[0]).strip()
                modo = 'val'
            elif la in ('observações', 'obs'):
                if vals:
                    cur['obs'].append(' '.join(str(v) for v in vals if isinstance(v, str)))
                modo = 'obs'
            elif la and (la.startswith('ajuste') or la == 'atualizada'):
                modo = None
            elif modo == 'modo' and isinstance(a, (int, float)) and len(r) > 1 and isinstance(r[1], str) and r[1].strip():
                cur['modo'].append(' '.join(r[1].split()))
            elif modo == 'armaz' and vals and not cur['armaz']:
                cur['armaz'] = str(vals[0]).strip()
            elif modo == 'val' and vals and not cur['validade']:
                cur['validade'] = str(vals[0]).strip()
            elif modo == 'obs':
                t = ' '.join(str(v) for v in ([a] if isinstance(a, str) else []) + vals if isinstance(v, str)).strip()
                if t and low(t) not in LABELS:
                    cur['obs'].append(t)
            elif modo is None and isinstance(a, int) and len(r) > 2 and isinstance(r[1], str) and r[1].strip() and not cur['modo']:
                cur['itens'].append((' '.join(r[1].split()), r[2]))
    return out


def numero(v):
    if isinstance(v, (int, float)):
        return float(v)
    m = re.fullmatch(r'\s*([\d.,]+)\s*(kg|kgs)?\s*', str(v or ''), re.I)
    return float(m.group(1).replace(',', '.')) if m else None


def casar(impressas, fichas):
    fichas = [f for f in fichas if f['tipo'] == 'preparo']
    por_nome = {simples(f['nome']): f for f in fichas}
    por_cod = {f['cod']: f for f in fichas}
    casadas = {}
    for s in impressas:
        t = simples(s['titulo'])
        if t in IMPRESSA:
            f = por_cod.get(IMPRESSA[t])
        else:
            f = por_nome.get(t)
            if not f:
                base = re.sub(r' (corte|empanamento|boleamento)$', '', t)
                f = por_nome.get(base) or next((g for n, g in por_nome.items() if n.startswith(base) or base.startswith(n + ' ')), None)
            if not f and s['codigo'] in por_cod and set(t.split()) & set(simples(por_cod[s['codigo']]['nome']).split()) - {'de', 'com', 'e'}:
                f = por_cod[s['codigo']]
        if f:
            casadas.setdefault(f['cod'], []).append(s)
    return casadas


def texto_impresso(secoes):
    """Junta as etapas (corte, empanamento…) de uma ficha em um texto só."""
    modo, varias = [], len(secoes) > 1
    for s in secoes:
        if not s['modo']:
            continue
        if varias:
            etapa = re.sub(r'^.*?(CORTE|EMPANAMENTO|BOLEAMENTO)$', r'\1', s['titulo'].upper())
            modo.append(f'{etapa.capitalize()}:' if etapa != s['titulo'].upper() else s['titulo'].capitalize() + ':')
        modo += s['modo']
    pega = lambda k: next((s[k] for s in secoes if s[k]), None)  # noqa: E731
    val = pega('validade')
    dias = re.search(r'(\d+)\s*dia', val or '')
    s0 = next((s for s in secoes if s['porcoes'] and numero(s['rend'])), None)
    porcao = None
    if s0 and isinstance(s0['porcoes'], (int, float)) and s0['porcoes'] > 0:
        porcao = (s0['porcao_nome'], round(numero(s0['rend']) / s0['porcoes'], 4))
    obs = [o for s in secoes for o in s['obs'] if o]
    if val and not dias:
        obs.append('Validade: ' + val)
    return {
        'modo': '\n'.join(modo) or None, 'conservacao': pega('armaz'), 'validade': int(dias.group(1)) if dias else None,
        'obs': '\n'.join(dict.fromkeys(obs)) or None, 'resp': pega('resp'), 'porcao': porcao,
    }


def main():
    fts, impressas_path, produtos_path, relatorio = sys.argv[1:5]
    fichas = ler_fts(fts)
    pasta = produtos_path.rsplit('/', 1)[0]
    unid_mat = {r[0]: r[3] for r in openpyxl.load_workbook(pasta + '/materiais.xlsx', read_only=True, data_only=True).active.iter_rows(values_only=True)
                if isinstance(r[0], int)}
    codigos = {f['cod'] for f in fichas}
    prods = {r[0]: r for r in openpyxl.load_workbook(produtos_path, read_only=True, data_only=True).active.iter_rows(values_only=True)
             if isinstance(r[0], int)}
    for f in fichas:
        f['tipo'] = 'produto' if f['cod'] in prods else 'preparo'
    impressas = ler_impressas(impressas_path)
    casadas = casar(impressas, fichas)
    alias = {v: k for k, v in ALIAS.items()}

    linhas_f, linhas_i = [], []
    vistos = {}
    for f in fichas:
        tipo = f['tipo']
        nome = nome_bom(prods[f['cod']][1] if tipo == 'produto' else f['nome'])
        chave = (simples(nome), tipo)
        if chave in vistos:
            nome = f'{nome} ({f["cod"]})'
        vistos[chave] = 1
        t = texto_impresso(casadas.get(f['cod'], []))
        linhas_f.append((str(f['cod']), nome, tipo, not INATIVA.search(f['nome']), alias.get(f['cod']), t['modo'], t['conservacao'],
                         t['validade'], t['obs'], t['resp'], t['porcao'][0] if t['porcao'] else None,
                         t['porcao'][1] if t['porcao'] else None))
        for k, it in enumerate(f['itens']):
            if it['cod'] == f['cod']:
                continue
            linhas_i.append((str(f['cod']), k + 1, str(it['cod']), it['qtd'], it['aprov'], it['delivery'], it['cod'] in codigos))

    print('-- Fichas técnicas da Eclética (Heitor, 10/10), geradas por scripts/importar_fichas_ecletica.py.')
    print('-- Rendimento 1 (1 kg, 1 L ou 1 unidade): é como a Eclética guarda. A ficha impressa multiplica.')
    print("""create function pg_temp.norm(t text) returns text language sql immutable as $$
  select regexp_replace(lower(translate(trim(t), 'ÁÀÂÃÄÉÈÊËÍÌÎÏÓÒÔÕÖÚÙÛÜÇáàâãäéèêëíìîïóòôõöúùûüç', 'AAAAAEEEEIIIIOOOOOUUUUCaaaaaeeeeiiiiooooouuuuc')), '\\s+', ' ', 'g')
$$;""")
    valores('_f', 'codigo text primary key, nome text, tipo text, ativo boolean, nome_evento text, modo text, conservacao text, '
            'validade int, observacoes text, responsavel text, porcao_nome text, porcao_qtd numeric', linhas_f)
    valores('_fi', 'ficha text, ordem int, item text, quantidade numeric, aproveitamento numeric, so_delivery boolean, tem_ficha boolean', linhas_i)
    print(open(__file__.rsplit('/', 1)[0] + '/importar_fichas_ecletica.sql').read())

    # Relatório: o que a ficha impressa tem de diferente da Eclética.
    with open(relatorio, 'w') as out:
        out.write('# Fichas impressas × Eclética (10/10)\n\n')
        out.write('O portal usa as quantidades da Eclética (por 1 kg). Abaixo, onde a ficha impressa de hoje diz outra coisa.\n')
        out.write('Quantidades comparadas por kg (ou litro) pronto quando a impressa tem o rendimento em número.\n\n')
        por_cod = {f['cod']: f for f in fichas}
        sem = [s['titulo'] for s in impressas if not any(s in v for v in casadas.values())]
        for cod, secoes in sorted(casadas.items(), key=lambda x: por_cod[x[0]]['nome']):
            f = por_cod[cod]
            ecl = {simples(i['nome']): i for i in f['itens']}
            difs = []
            for s in secoes:
                rend = numero(s['rend']) if unid_mat.get(cod) in ('Kilos', 'Litros') else None
                for nome, q in s['itens']:
                    n = simples(nome)
                    achado = ecl.get(n) or next((v for k, v in ecl.items() if n and (n in k or k in n)), None)
                    if not achado:
                        difs.append(f'- **{nome}** está na impressa e não na Eclética')
                    elif rend and isinstance(q, (int, float)) and q > 0:
                        porkg = q / rend
                        if abs(porkg - achado['qtd']) > max(0.002, 0.1 * achado['qtd']):
                            difs.append(f'- {nome}: impressa {porkg:.3f} por kg, Eclética {achado["qtd"]:.3f}'.replace('.', ','))
                for k, v in ecl.items():
                    if not any(simples(nome) and (simples(nome) in k or k in simples(nome)) for nome, _ in s['itens']):
                        difs.append(f'- {v["nome"]} está na Eclética e não na impressa')
            if difs:
                out.write(f'## {nome_bom(f["nome"])} ({cod})\n' + '\n'.join(dict.fromkeys(difs)) + '\n\n')
        if sem:
            out.write('## Impressas sem fórmula na Eclética\n' + '\n'.join(f'- {t}' for t in sem) + '\n')
    print(f'-- {len(linhas_f)} fichas, {len(linhas_i)} itens', file=sys.stderr)
    print(f'-- impressas casadas: {sum(len(v) for v in casadas.values())} de {len(impressas)}', file=sys.stderr)


if __name__ == '__main__':
    main()
