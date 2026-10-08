"""Gera a migration com insumos, fornecedores e fichas da planilha DPEN 25-26 (enviada pelo Heitor em 08/10/2026).

Uso: python3 -I scripts/importar_dpen.py <planilha.xlsx> > supabase/migrations/0035_dados_planilha_dpen.sql
Regras (as mesmas que a tela usa para o custo):
- quantidade da ficha está na unidade do insumo; comprar = quantidade ÷ aproveitamento;
- pré-preparo vira ficha do tipo 'preparo' (rendimento 1 na sua unidade) e entra nas outras fichas como sub-ficha;
- insumos de 'Produção própria' que têm ficha não viram insumo (o custo vem da ficha);
- preço de venda padrão = último preço médio vendido no DPEN.
"""
import collections, datetime, sys
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
num = lambda v: v if isinstance(v, (int, float)) and not isinstance(v, bool) else None
q = lambda s: 'null' if s is None else "'" + str(s).replace("'", "''") + "'"
n = lambda v: 'null' if v is None else repr(round(float(v), 4))
chave = lambda s: s.strip().lower()

mats = [r for r in list(wb['Lista de Materiais'].iter_rows(values_only=True))[1:] if r[0]]
fich = [r for r in list(wb['Fichas técnicas'].iter_rows(values_only=True))[1:] if r[3]]

ultimo = {}
for r in list(wb['Produtos por Dia(DB)'].iter_rows(values_only=True))[2:]:
    if not r[0] or not isinstance(r[3], datetime.datetime) or num(r[6]) is None:
        continue
    k = r[5].strip().upper()
    if k not in ultimo or r[3] > ultimo[k][0]:
        ultimo[k] = (r[3], r[6])

# Fichas: (tipo, nome) -> dados e itens
fichas = collections.OrderedDict()
for r in fich:
    tipo = 'preparo' if r[0] == 'Pré Preparo' else 'produto'
    k = (tipo, chave(r[3]))
    f = fichas.setdefault(k, {'nome': r[3].strip(), 'tipo': tipo, 'linha': None, 'ativo': r[1] == 'ATIVO', 'itens': []})
    linha = r[4] if isinstance(r[4], str) and not r[4].startswith('=') else None
    f['linha'] = f['linha'] or linha
    apr = num(r[5])
    f['itens'].append({'nome': r[2].strip(), 'qtd': num(r[6]), 'apr': apr if apr and 0 < apr <= 1 else 1})
for f in fichas.values():
    if f['linha'] is None and f['tipo'] == 'produto':
        f['linha'] = 'Bebidas'
preparos = {k[1] for k in fichas if k[0] == 'preparo'}

GENERICOS = {None, '', 'Produção própria', 'Fornecedor mais barato', 'Fornecedor'}
UNIDADE = {'kg': ('kg', None), 'un': ('un', None), 'fardos': ('un', 'fardo'), 'pote': ('kg', 'pote'), 'baldes': ('kg', 'balde'), 'cx': ('kg', 'caixa')}

insumos = collections.OrderedDict()
for r in mats:
    k = chave(r[0])
    if k in preparos:
        continue
    un, emb = UNIDADE.get(r[7], ('kg', None))
    preco = num(r[11])
    obs = None
    if r[7] in ('pote', 'baldes', 'cx'):
        obs = f'Na planilha a unidade de compra é "{r[7]}"; conferir se o preço é por kg.'
    if r[8] == 'Produção própria':
        obs = 'Produção própria, sem ficha na planilha.'
    insumos[k] = {'nome': r[0].strip(), 'categoria': r[1], 'unidade': un, 'embalagem': emb, 'preco': preco or None,
                  'fornecedor': None if r[8] in GENERICOS else r[8].strip(), 'ativo': r[2] == 'ATIVO', 'obs': obs}
for f in fichas.values():
    for i in f['itens']:
        k = chave(i['nome'])
        if k not in preparos and k not in insumos:
            insumos[k] = {'nome': i['nome'], 'categoria': None, 'unidade': 'un' if i['qtd'] == 1 else 'kg', 'embalagem': None,
                          'preco': None, 'fornecedor': None, 'ativo': False, 'obs': 'Criado a partir das fichas da planilha; sem preço.'}

# Custo pela mesma regra da tela.
memo = {}
def custo(tipo, k, pilha=()):
    if (tipo, k) in memo:
        return memo[(tipo, k)]
    total = 0.0
    for i in fichas[(tipo, k)]['itens']:
        ki = chave(i['nome'])
        if ki in preparos and ki != k and ki not in pilha:
            unit = custo('preparo', ki, pilha + (k,))
        else:
            unit = (insumos.get(ki) or {}).get('preco')
        if unit is None or i['qtd'] is None:
            continue
        total += i['qtd'] / i['apr'] * unit
    memo[(tipo, k)] = total
    return total

out = ['-- Gerado por scripts/importar_dpen.py a partir da planilha "DPEN 25-26.xlsx" (Heitor, 08/10/2026).',
       '-- Insumos, fornecedores e fichas de eventos. Ponto de partida: tudo pode ser corrigido pela tela (cria versão nova).',
       "select set_config('portal.origem_preco', 'planilha DPEN', true);", '']
fornecedores = sorted({v['fornecedor'] for v in insumos.values() if v['fornecedor']}, key=str.lower)
out.append('insert into fornecedores (nome) values\n  ' + ',\n  '.join(f'({q(x)})' for x in fornecedores) + ';\n')
linhas = []
for v in insumos.values():
    forn = f"(select id from fornecedores where lower(nome) = lower({q(v['fornecedor'])}))" if v['fornecedor'] else 'null'
    linhas.append(f"({q(v['nome'])}, {q(v['categoria'])}, {q(v['unidade'])}, {q(v['embalagem'])}, {n(v['preco'])}, {forn}, {q(v['obs'])}, {str(v['ativo']).lower()})")
out.append('insert into insumos (nome, categoria, unidade, embalagem, preco, fornecedor_id, observacao, ativo) values\n  ' + ',\n  '.join(linhas) + ';\n')

linhas = []
for (tipo, k), f in fichas.items():
    unidade = 'un' if tipo == 'produto' else ({'kg': 'kg', 'un': 'un'}.get(next((r[7] for r in mats if chave(r[0]) == k), 'kg'), 'kg'))
    origem = 'revenda' if f['linha'] == 'Bebidas' and len(f['itens']) == 1 else 'propria'
    preco = ultimo.get(f['nome'].upper(), (None, None))[1] if tipo == 'produto' else None
    linhas.append(f"({q(f['nome'])}, {q(tipo)}, {q(f['linha'])}, {q(origem)}, {q(unidade)}, {n(preco)}, {str(f['ativo']).lower()}, 1)")
out.append('insert into receitas (nome, tipo, linha, origem, unidade, preco_venda, ativo, versao_atual) values\n  ' + ',\n  '.join(linhas) + ';\n')

nota = 'Importada da planilha DPEN 25-26 em 08/10/2026.'
linhas = []
for (tipo, k), f in fichas.items():
    linhas.append(f"((select id from receitas where lower(nome) = {q(k)} and tipo = {q(tipo)}), 1, 1, {n(custo(tipo, k))}, {q(nota)})")
out.append('insert into receita_versoes (receita_id, numero, rendimento, custo_total, nota) values\n  ' + ',\n  '.join(linhas) + ';\n')

linhas = []
for (tipo, k), f in fichas.items():
    versao = f"(select v.id from receita_versoes v join receitas r on r.id = v.receita_id where lower(r.nome) = {q(k)} and r.tipo = {q(tipo)} and v.numero = 1)"
    for o, i in enumerate(f['itens'], 1):
        if i['qtd'] is None or i['qtd'] <= 0:
            continue
        ki = chave(i['nome'])
        if ki in preparos and ki != k:
            ins, sub = 'null', f"(select id from receitas where lower(nome) = {q(ki)} and tipo = 'preparo')"
        else:
            ins, sub = f"(select id from insumos where lower(nome) = {q(ki)})", 'null'
        linhas.append(f"({versao}, {o}, {ins}, {sub}, {n(i['qtd'])}, {n(i['apr'])})")
out.append('insert into receita_itens (versao_id, ordem, insumo_id, sub_receita_id, quantidade, aproveitamento) values\n  ' + ',\n  '.join(linhas) + ';')
print('\n'.join(out))
print(f'-- {len(fornecedores)} fornecedores, {len(insumos)} insumos, {len(fichas)} fichas, {len(linhas)} itens', file=sys.stderr)
for (tipo, k), f in fichas.items():
    if f['ativo']:
        print(f"{tipo:8} {f['nome']:28} custo {custo(tipo, k):8.2f}", file=sys.stderr)
