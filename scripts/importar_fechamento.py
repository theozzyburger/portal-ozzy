"""Gera os dados da migration 0059 a partir da planilha "LISTA FECHAMENTO UN1.xlsx" (Heitor, 09/10).

Uso: python3 -I scripts/importar_fechamento.py <planilha.xlsx> > supabase/migrations/0059_dados_fechamento.sql

- Cozinha: materiais da Matriz com tipo 22, mais os 5 códigos novos que aparecem nas fotos do fechamento de 08/10.
- Atendimento: materiais da Matriz com tipo 21.
- As duas lojas (PSD = UN1, Vila = UN2) têm as mesmas listas; o estoque ideal de cada dia vem de "Dados UN1" e
  "Dados UN2" (0 = segunda). Item sem ideal fica sem sugestão de pedido.
- A unidade de contagem é a da Eclética (Kg, Uni, Lts, Pct, Cxs, GL); o item de estoque fica em kg, l ou un.
- Nome em maiúsculas vira só a primeira letra maiúscula.
"""
import sys, unicodedata
import openpyxl

TIPO_SETOR = {22: 'cozinha', 21: 'atendimento'}
UNID = {0: 'Kg', 1: 'Lts', 2: 'Uni', 3: 'Fds', 4: 'Pct', 5: 'Cxs', 6: 'Pés', 7: 'Pçs', 8: 'Bis', 9: 'GL'}
INSUMO = {'Kg': 'kg', 'Lts': 'l'}
DIAS = ['Segunda', 'Terça', 'Quarta', 'Quinta', 'Sexta', 'Sábado', 'Domingo']
# Fora da Matriz, vistos nas fotos de 08/10 (cozinha da UN1).
NOVOS = [
    ('100619', 'Água sanitária 2L', 'GL', False),
    ('100627', 'Asinha de frango', 'Kg', False),
    ('100624', 'Cogumelo Paris champetit', 'Kg', False),
    ('100617', 'Maionese defumada juicy', 'Kg', True),
    ('100618', 'Maionese salsa de trufas', 'Kg', True),
]

# Marcados como pré-preparo na Eclética por engano.
NAO_PREPARO = {'100328'}  # Bacias de Inox


def nome_bom(s):
    s = ' '.join(str(s).split())
    # Palavras só de letras com 3 ou mais (ignora "200ml", "c/100"): todas em maiúsculas = nome gritado.
    palavras = [w for w in s.split() if w.isalpha() and len(w) >= 3]
    if palavras and all(w.isupper() for w in palavras):
        s = s[0].upper() + s[1:].lower()
    return s


def chave(s):
    return unicodedata.normalize('NFD', s.lower()).encode('ascii', 'ignore').decode()


def sql(v):
    if v is None:
        return 'null'
    if isinstance(v, bool):
        return 'true' if v else 'false'
    if isinstance(v, (int, float)):
        return repr(round(v, 3)).rstrip('0').rstrip('.') if isinstance(v, float) else str(v)
    return "'" + str(v).replace("'", "''") + "'"


def ideais(wb, aba):
    r = {}
    for row in list(wb[aba].iter_rows(values_only=True))[1:]:
        if row[1] is None or row[2] not in DIAS:
            continue
        cod = str(row[1]).strip()
        q = row[4] if isinstance(row[4], (int, float)) else None
        r.setdefault(cod, [None] * 7)[DIAS.index(row[2])] = q
    return r


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    itens = []  # (codigo, nome, unidade_contagem, pre_preparo, setor)
    for row in list(wb['Matriz'].iter_rows(values_only=True))[1:]:
        if row[0] is None or row[3] not in TIPO_SETOR:
            continue
        cod = str(row[0]).strip()
        itens.append((cod, nome_bom(row[1]), UNID.get(row[2], 'Uni'), row[11] == 'S' and cod not in NAO_PREPARO, TIPO_SETOR[row[3]]))
    itens += [(c, n, u, p, 'cozinha') for c, n, u, p in NOVOS]

    # Nomes repetidos (o cadastro de insumos não aceita dois iguais): o segundo leva o código.
    vistos = set()
    unicos = []
    for c, n, u, p, s in itens:
        if chave(n) in vistos:
            n = f'{n} ({c})'
        vistos.add(chave(n))
        unicos.append((c, n, u, p, s))
    itens = sorted(unicos, key=lambda i: (i[4], chave(i[1])))

    un1, un2 = ideais(wb, 'Dados UN1'), ideais(wb, 'Dados UN2')
    arr = lambda v: 'null' if v is None or all(x is None for x in v) else "array[" + ', '.join(sql(x) for x in v) + "]::numeric[]"

    print('-- Gerado por scripts/importar_fechamento.py a partir de "LISTA FECHAMENTO UN1.xlsx" (09/10).')
    print(f'-- {sum(1 for i in itens if i[4] == "cozinha")} itens da cozinha e {sum(1 for i in itens if i[4] == "atendimento")} do atendimento, nas duas lojas.')
    print('create temp table _fech (codigo text, nome text, unidade_contagem text, pre_preparo boolean, setor text, ordem int,')
    print('  ideal_psd numeric[], ideal_va numeric[]) on commit drop;')
    print('insert into _fech values')
    linhas = []
    ordem = {'cozinha': 0, 'atendimento': 0}
    for c, n, u, p, s in itens:
        ordem[s] += 1
        linhas.append(f'  ({sql(c)}, {sql(n)}, {sql(u)}, {sql(p)}, {sql(s)}, {ordem[s]}, {arr(un1.get(c))}, {arr(un2.get(c))})')
    print(',\n'.join(linhas) + ';')
    print('''
-- Liga aos insumos que já existem pelo nome; os outros entram no cadastro.
update insumos i set ecletica_codigo = f.codigo
from _fech f
where i.ecletica_codigo is null and lower(i.nome) = lower(f.nome)
  and not exists (select 1 from insumos x where x.ecletica_codigo = f.codigo);
insert into insumos (nome, categoria, unidade, ecletica_codigo, observacao)
select f.nome, case f.setor when 'cozinha' then 'Cozinha das lojas' else 'Atendimento das lojas' end,
  case f.unidade_contagem when 'Kg' then 'kg' when 'Lts' then 'l' else 'un' end, f.codigo,
  'Lista de fechamento das lojas (código da Eclética).'
from _fech f
where not exists (select 1 from insumos x where x.ecletica_codigo = f.codigo)
  and not exists (select 1 from insumos x where lower(x.nome) = lower(f.nome));
update insumos i set pre_preparo = true from _fech f where i.ecletica_codigo = f.codigo and f.pre_preparo;

insert into fechamento_itens (unidade_id, setor, insumo_id, unidade_contagem, ordem, ideal)
select u.id, f.setor, i.id, f.unidade_contagem, f.ordem, case u.id when 'burger-psd' then f.ideal_psd else f.ideal_va end
from _fech f
join insumos i on i.ecletica_codigo = f.codigo
cross join (values ('burger-psd'), ('burger-va')) u (id)
on conflict (unidade_id, setor, insumo_id) do nothing;''')


if __name__ == '__main__':
    main(sys.argv[1])
