"""Gera a migration 0039 a partir da planilha "Fichas de preparo.xlsx" (enviada pelo Heitor em 08/10).

Uso: python3 -I scripts/importar_fichas_preparo.py <planilha.xlsx> > supabase/migrations/0039_fichas_de_preparo.sql

Cada aba é um preparo: insumos, rendimento, modo de preparo, armazenamento e validade.
- Preparo que já existe no catálogo (vindo das fichas técnicas DPEN): ganha modo de preparo, armazenamento e
  validade. Os insumos continuam os da ficha técnica (o Heitor disse "usa uma ou outra").
- Preparo que não existe: entra no catálogo sem ficha (versão 0), com o modo de preparo e a lista de insumos da
  planilha no texto, para a gestão montar a ficha com os insumos certos do cadastro.
- A aba "Blood Lemonade" é cópia do pesto e fica de fora.
"""
import re, sys
import openpyxl

# Aba da planilha -> preparo que já está no catálogo.
EXISTENTES = {
    'Prato de pizza': 'Prato de pizza',
    'Molho Branco': 'Molho Branco',
    'Polpettone': 'Polpettone',
    'Coleslaw': 'Coleslaw',
    'Pesto de Manjericão': 'Pesto de manjericão',
    'Calda de frutas vermelhas': 'Calda de Frutas vermelhas caseira',
    'Crema de Queijo': 'Crema de Queijo',
    'Massa Cheese Cake': 'Massa de Cheesecake',
    'Base Cheese Cake': 'Base Amenteigada de Cheesecake',
}
# Aba -> nome do preparo novo (e a unidade do que rende).
NOVOS = {
    'Molho de tomate': ('Molho de tomate', 'kg'),
    'Fettuccine': ('Fettuccine pré-cozido', 'kg'),
    'Mortadela Fatiada': ('Mortadela fatiada', 'kg'),
    'Mussarela Ralada': ('Mussarela ralada', 'kg'),
    'Mussarela Fatiada': ('Mussarela fatiada', 'kg'),
    'Bacon em cubos na chapa': ('Bacon em cubos na chapa', 'kg'),
    'Parma Pacote': ('Parma em pacote', 'kg'),
    'Parma e Copa': ('Parma e copa', 'kg'),
    'Purê de batata pronto': ('Purê de batata pronto', 'kg'),
    'Focaccia Forma': ('Focaccia em forma', 'un'),
    'Straciatella': ('Stracciatella', 'kg'),
    'Árabe': ('Recheio árabe (kafta e molho)', 'un'),
    'Pastrami': ('Pastrami fatiado', 'kg'),
    'Sal de alecrim': ('Sal de alecrim', 'kg'),
    'Mostarda da Casa': ('Mostarda da casa', 'kg'),
}
FORA = {'Blood Lemonade'}


def ler(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    for aba in wb.sheetnames:
        itens, passos, arm, val, rend, sec = [], [], None, None, None, None
        for r in wb[aba].iter_rows(values_only=True):
            v = [c for c in r if c is not None]
            if not v:
                continue
            t = str(v[0]).strip().upper() if isinstance(v[0], str) else None
            if t in ('INSUMOS', 'MODO DE PREPARO', 'OBS') or (t == 'ARMAZENAMENTO' and len(v) == 1):
                sec = t
            elif t == 'ARMAZENAMENTO':
                arm = str(v[1]).strip()
            elif t == 'VALIDADE':
                val = str(v[1]).strip() if len(v) > 1 else None
            elif t == 'RENDIMENTO':
                rend = v[1] if len(v) > 1 else None
            elif t == '#':
                continue
            elif sec == 'INSUMOS' and len(v) >= 3 and isinstance(v[1], str) and isinstance(v[2], (int, float)):
                itens.append((v[1].strip(), v[2]))
            elif sec == 'MODO DE PREPARO' and len(v) >= 2 and str(v[1]).strip().lower() != 'a definir':
                passos.append(str(v[1]).strip())
        yield aba, itens, passos, arm, val, rend


def dias(val):
    m = re.match(r'(\d+)', val or '')
    return int(m.group(1)) if m else None


def q(s):
    return 'null' if s is None else "'" + str(s).replace("'", "''") + "'"


def num(x):
    return f'{x:g}'.replace('.', ',')


def main():
    print('-- Fichas de preparo (planilha "Fichas de preparo.xlsx", enviada pelo Heitor em 08/10).')
    print('-- Gerado por scripts/importar_fichas_preparo.py. Só acrescenta a coluna modo_preparo, preenche texto e')
    print('-- cria preparos novos sem ficha (versão 0): os insumos e custos que já existem não mudam.')
    print()
    print('alter table receitas add column modo_preparo text;')
    print()
    vistos = set()
    for aba, itens, passos, arm, val, rend in ler(sys.argv[1]):
        if aba in FORA:
            continue
        modo = '\n'.join(f'{i}. {p}' for i, p in enumerate(passos, 1)) or None
        if aba in EXISTENTES:
            vistos.add(aba)
            print(f'update receitas set modo_preparo = coalesce(modo_preparo, {q(modo)}),')
            print(f'  conservacao = coalesce(conservacao, {q(arm)}), validade_dias = coalesce(validade_dias, {dias(val) if dias(val) is not None else "null"})')
            print(f"where lower(nome) = lower({q(EXISTENTES[aba])}) and tipo = 'preparo';")
        elif aba in NOVOS:
            vistos.add(aba)
            nome, unidade = NOVOS[aba]
            lista = 'Insumos na planilha (para 1 receita' + (f', rende {num(rend)}' if rend else '') + '): ' + '; '.join(f'{i} {num(x)}' for i, x in itens) + '.'
            texto = (modo + '\n\n' if modo else '') + lista
            print("insert into receitas (nome, tipo, origem, unidade, conservacao, validade_dias, modo_preparo, versao_atual)")
            print(f"select {q(nome)}, 'preparo', 'propria', '{unidade}', {q(arm)}, {dias(val) if dias(val) is not None else 'null'}, {q(texto)}, 0")
            print(f"where not exists (select 1 from receitas where lower(nome) = lower({q(nome)}) and tipo = 'preparo');")
        else:
            raise SystemExit(f'Aba sem destino: {aba}')
    faltam = (set(EXISTENTES) | set(NOVOS)) - vistos
    if faltam:
        raise SystemExit(f'Abas esperadas e não encontradas: {faltam}')


main()
